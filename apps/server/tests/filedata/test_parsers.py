from datetime import date
from pathlib import Path

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
import pytest

from datapulse.contracts.dataset import FileFormat
from datapulse.filedata.parsers import FileParseInvalid, excel_sheet_names, parse_file

pytestmark = pytest.mark.anyio


def write_excel(path: Path) -> None:
    workbook = openpyxl.Workbook()
    first = workbook.active
    first.title = "销售"
    first.append(["month", "amount"])
    first.append(["2026-01", 10])
    first.append(["2026-02", 20])
    second = workbook.create_sheet("库存")
    second.append(["sku", "stock"])
    second.append(["A-1", 5])
    workbook.save(path)


async def test_parse_csv_utf8_bom_and_gbk(tmp_path: Path) -> None:
    utf8_path = tmp_path / "sales-utf8.csv"
    utf8_path.write_text("region,amount\n华东,10\n华南,20\n", encoding="utf-8")
    bom_path = tmp_path / "sales-bom.csv"
    bom_path.write_bytes("region,amount\n华东,10\n".encode("utf-8-sig"))
    gbk_path = tmp_path / "sales-gbk.csv"
    gbk_path.write_bytes("region,amount\n华东,10\n".encode("gb18030"))

    utf8 = await parse_file(utf8_path, FileFormat.CSV, sheet_name=None, max_rows=2)
    bom = await parse_file(bom_path, FileFormat.CSV, sheet_name=None, max_rows=2)
    gbk = await parse_file(gbk_path, FileFormat.CSV, sheet_name=None, max_rows=2)

    assert utf8.row_count == 2
    assert utf8.fields[0].name == "region"
    assert {field.name: field.data_type for field in utf8.fields} == {
        "region": "string",
        "amount": "integer",
    }
    assert utf8.sample_rows[0]["region"] == "华东"
    assert utf8.sample_rows[0]["amount"] == 10
    assert bom.sample_rows[0]["region"] == "华东"
    assert gbk.sample_rows[0]["region"] == "华东"


async def test_parse_chinese_csv_infers_boolean_numeric_and_datetime_columns(
    tmp_path: Path,
) -> None:
    path = tmp_path / "中文销售.csv"
    path.write_text(
        "订单号,是否完成,金额,发生时间\n"
        "A-001,是,10.5,2026-09-22T08:30:00\n"
        "A-002,否,20,2026-09-23T08:30:00\n",
        encoding="utf-8",
    )

    parsed = await parse_file(path, FileFormat.CSV, sheet_name=None, max_rows=10)

    assert [(field.name, field.data_type.value) for field in parsed.fields] == [
        ("订单号", "string"),
        ("是否完成", "boolean"),
        ("金额", "number"),
        ("发生时间", "datetime"),
    ]
    assert parsed.sample_rows[0] == {
        "订单号": "A-001",
        "是否完成": True,
        "金额": 10.5,
        "发生时间": "2026-09-22T08:30:00",
    }


async def test_parse_excel_requires_sheet_selection_and_exports_normalized_path(
    tmp_path: Path,
) -> None:
    excel_path = tmp_path / "sales.xlsx"
    write_excel(excel_path)

    parsed = await parse_file(excel_path, FileFormat.EXCEL, sheet_name="库存", max_rows=10)

    assert parsed.row_count == 1
    assert parsed.fields[0].name == "sku"
    assert parsed.sample_rows[0]["sku"] == "A-1"
    assert parsed.normalized_path.suffix == ".parquet"
    assert parsed.normalized_path.is_file()


async def test_excel_sheet_names_preserve_order_and_select_distinct_sheets(
    tmp_path: Path,
) -> None:
    excel_path = tmp_path / "workbook.xlsx"
    workbook = openpyxl.Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["metric", "value"])
    summary.append(["revenue", 100])
    detail = workbook.create_sheet("Detail")
    detail.append(["order_id", "amount"])
    detail.append(["A-1", 60])
    workbook.save(excel_path)

    names = await excel_sheet_names(excel_path)
    summary_result = await parse_file(
        excel_path,
        FileFormat.EXCEL,
        sheet_name="Summary",
        max_rows=10,
    )
    detail_result = await parse_file(
        excel_path,
        FileFormat.EXCEL,
        sheet_name="Detail",
        max_rows=10,
    )

    assert names == ("Summary", "Detail")
    assert summary_result.fields[0].name == "metric"
    assert summary_result.sample_rows[0]["metric"] == "revenue"
    assert detail_result.fields[0].name == "order_id"
    assert detail_result.sample_rows[0]["order_id"] == "A-1"


async def test_parse_json_requires_object_array_root_and_rejects_nested_objects(
    tmp_path: Path,
) -> None:
    valid_path = tmp_path / "sales.json"
    valid_path.write_text(
        '[{"region":"east","amount":10},{"region":"west","amount":20}]',
        encoding="utf-8",
    )
    empty_path = tmp_path / "empty.json"
    empty_path.write_text("[]", encoding="utf-8")
    nested_path = tmp_path / "nested.json"
    nested_path.write_text(
        '[{"region":"east","detail":{"amount":10}}]',
        encoding="utf-8",
    )

    parsed = await parse_file(valid_path, FileFormat.JSON, sheet_name=None, max_rows=10)
    assert parsed.row_count == 2
    assert parsed.sample_rows[0]["amount"] == 10

    with pytest.raises(FileParseInvalid):
        await parse_file(empty_path, FileFormat.JSON, sheet_name=None, max_rows=10)
    with pytest.raises(FileParseInvalid):
        await parse_file(nested_path, FileFormat.JSON, sheet_name=None, max_rows=10)


async def test_parse_parquet_maps_numeric_and_date_fields(tmp_path: Path) -> None:
    parquet_path = tmp_path / "sales.parquet"
    table = pa.table(
        {
            "day": pa.array([date(2026, 1, 1), date(2026, 1, 2)], type=pa.date32()),
            "amount": pa.array([10.5, 12.0], type=pa.float64()),
            "count": pa.array([1, 2], type=pa.int64()),
        }
    )
    pq.write_table(table, parquet_path)

    parsed = await parse_file(parquet_path, FileFormat.PARQUET, sheet_name=None, max_rows=10)

    types_by_field = {field.name: field.data_type for field in parsed.fields}
    assert parsed.row_count == 2
    assert types_by_field["day"] == "date"
    assert types_by_field["amount"] == "number"
    assert types_by_field["count"] == "integer"


async def test_file_formats_share_scalar_inference_semantics(tmp_path: Path) -> None:
    json_path = tmp_path / "typed.json"
    json_path.write_text(
        '[{"amount":"10","enabled":"是","day":"2026-09-22"}]',
        encoding="utf-8",
    )
    excel_path = tmp_path / "typed.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["amount", "enabled", "day"])
    sheet.append(["10", "是", "2026-09-22"])
    workbook.save(excel_path)

    json_result = await parse_file(json_path, FileFormat.JSON, sheet_name=None, max_rows=10)
    excel_result = await parse_file(excel_path, FileFormat.EXCEL, sheet_name=None, max_rows=10)

    expected = {"amount": "integer", "enabled": "boolean", "day": "date"}
    assert {field.name: field.data_type.value for field in json_result.fields} == expected
    assert {field.name: field.data_type.value for field in excel_result.fields} == expected
    assert json_result.sample_rows[0] == excel_result.sample_rows[0] == {
        "amount": 10,
        "enabled": True,
        "day": "2026-09-22",
    }


async def test_all_file_formats_share_scalar_inference_semantics(tmp_path: Path) -> None:
    csv_path = tmp_path / "typed.csv"
    csv_path.write_text(
        "amount,enabled,day\n10,是,2026-09-22\n",
        encoding="utf-8",
    )
    json_path = tmp_path / "typed.json"
    json_path.write_text(
        '[{"amount":"10","enabled":"是","day":"2026-09-22"}]',
        encoding="utf-8",
    )
    excel_path = tmp_path / "typed.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["amount", "enabled", "day"])
    sheet.append(["10", "是", "2026-09-22"])
    workbook.save(excel_path)
    parquet_path = tmp_path / "typed.parquet"
    pq.write_table(
        pa.table(
            {
                "amount": ["10"],
                "enabled": ["是"],
                "day": ["2026-09-22"],
            }
        ),
        parquet_path,
    )

    results = []
    for path, file_format in (
        (csv_path, FileFormat.CSV),
        (excel_path, FileFormat.EXCEL),
        (json_path, FileFormat.JSON),
        (parquet_path, FileFormat.PARQUET),
    ):
        results.append(await parse_file(path, file_format, sheet_name=None, max_rows=10))

    expected = {"amount": "integer", "enabled": "boolean", "day": "date"}
    assert all(
        {field.name: field.data_type.value for field in result.fields} == expected
        for result in results
    )


async def test_parse_file_limits_sample_rows_without_losing_total_row_count(tmp_path: Path) -> None:
    csv_path = tmp_path / "sales.csv"
    csv_path.write_text("region,amount\nnorth,10\nsouth,20\nwest,30\n", encoding="utf-8")

    parsed = await parse_file(csv_path, FileFormat.CSV, sheet_name=None, max_rows=2)

    assert parsed.row_count == 3
    assert len(parsed.sample_rows) == 2
