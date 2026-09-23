import asyncio
import csv
import io
import json
import os
import shutil
import stat
from collections import OrderedDict
from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
from weakref import WeakKeyDictionary

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
from pydantic import BaseModel, ConfigDict

from datapulse.contracts.common import JsonValue
from datapulse.contracts.dataset import DatasetField, DataType, FileFormat
from datapulse.errors import DataPulseError
from datapulse.filedata.inference import coerce_rows, fields_from_rows, infer_type, merge_types
from datapulse.operations.processing import (
    ProcessorBusy,
    ProcessorFailed,
    complete_cleanup,
    run_json_worker,
)


class FileParseInvalid(ValueError):
    code = "FILE_PARSE_INVALID"


class FileParseBusy(DataPulseError):
    def __init__(self) -> None:
        super().__init__("FILE_PROCESSING_BUSY", "File processing is busy. Retry shortly.", 429)


class ParsedFile(BaseModel):
    model_config = ConfigDict(frozen=True)

    fields: tuple[DatasetField, ...]
    row_count: int
    sample_rows: tuple[dict[str, JsonValue], ...]
    normalized_path: Path


_CACHE_MAX_ENTRIES = 64
_CACHE_MAX_BYTES = 8 * 1024 * 1024
type _Fingerprint = tuple[int, int, int, int, int]
type _ParseKey = tuple[Path, _Fingerprint, FileFormat, str | None, int]


@dataclass
class _ParseFlight:
    task: asyncio.Task[bytes]
    waiters: int = 0


@dataclass
class _ParseCache:
    entries: OrderedDict[_ParseKey, tuple[bytes, Path, _Fingerprint]] = field(
        default_factory=OrderedDict
    )
    flights: dict[_ParseKey, _ParseFlight] = field(default_factory=dict)
    size_bytes: int = 0

    def remove(self, key: _ParseKey) -> None:
        entry = self.entries.pop(key, None)
        if entry is not None:
            self.size_bytes -= len(entry[0])


# Completed entries contain bytes and filesystem metadata only. No task, lock,
# callback or cache entry keeps its event loop alive after pending work finishes.
_parse_caches: WeakKeyDictionary[asyncio.AbstractEventLoop, _ParseCache] = WeakKeyDictionary()


def _fingerprint(path: Path) -> _Fingerprint:
    info = path.lstat()
    if not stat.S_ISREG(info.st_mode):
        raise FileParseInvalid("The file must be a regular file.")
    return (info.st_dev, info.st_ino, info.st_size, info.st_mtime_ns, info.st_ctime_ns)


async def _parse_and_cache(cache: _ParseCache, key: _ParseKey) -> bytes:
    path, fingerprint, file_format, sheet_name, max_rows = key
    parsed = await _parse_file_uncached(
        path, file_format, sheet_name=sheet_name, max_rows=max_rows, fingerprint=fingerprint
    )
    if _fingerprint(path) != fingerprint:
        raise FileParseInvalid("The source changed during file processing.")
    normalized_fingerprint = _fingerprint(parsed.normalized_path)
    payload = parsed.model_dump_json().encode()
    if len(payload) <= _CACHE_MAX_BYTES:
        cache.remove(key)
        cache.entries[key] = (payload, parsed.normalized_path, normalized_fingerprint)
        cache.size_bytes += len(payload)
        while len(cache.entries) > _CACHE_MAX_ENTRIES or cache.size_bytes > _CACHE_MAX_BYTES:
            cache.remove(next(iter(cache.entries)))
    return payload


def _merge_type(current: DataType | None, value: object) -> DataType:
    return merge_types(current, infer_type(value))


def _normalize_value(value: object) -> JsonValue:
    if value is None:
        return None
    if isinstance(value, (bool, int, float, str)):
        return value
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raise FileParseInvalid("The file contains unsupported nested values.")


def _fields_from_rows(rows: Iterable[dict[str, object]]) -> tuple[DatasetField, ...]:
    return fields_from_rows(rows)


def _write_parquet(rows: list[dict[str, object]], target: Path) -> Path:
    target.parent.mkdir(parents=True, exist_ok=True)
    pq.write_table(pa.Table.from_pylist(rows), target)
    return target


def _detect_text(data: bytes) -> tuple[str, str]:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise FileParseInvalid("The file encoding is not supported.")


def _parse_csv(path: Path, *, max_rows: int) -> ParsedFile:
    raw = path.read_bytes()
    text, encoding = _detect_text(raw)
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise FileParseInvalid("The CSV header is invalid.")
    rows: list[dict[str, object]] = []
    for raw_row in reader:
        row = {key: value for key, value in raw_row.items() if key}
        rows.append(row)
    rows = coerce_rows(rows)
    fields = fields_from_rows(rows)
    normalized_path = path
    if encoding not in {"utf-8", "utf-8-sig"}:
        normalized_path = path.with_suffix(".normalized.csv")
        normalized_path.write_text(text, encoding="utf-8")
    return ParsedFile(
        fields=fields,
        row_count=len(rows),
        sample_rows=tuple(
            {key: _normalize_value(value) for key, value in row.items()} for row in rows[:max_rows]
        ),
        normalized_path=normalized_path,
    )


def _parse_excel(path: Path, *, sheet_name: str | None, max_rows: int) -> ParsedFile:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name is not None and sheet_name not in workbook.sheetnames:
            raise FileParseInvalid("The Excel sheet does not exist.")
        worksheet = workbook[sheet_name] if sheet_name is not None else workbook.active
        rows_iter = worksheet.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration as error:
            raise FileParseInvalid("The Excel sheet is empty.") from error
        headers = [str(cell).strip() if cell is not None else "" for cell in header_row]
        if not any(headers) or any(not header for header in headers):
            raise FileParseInvalid("The Excel header is invalid.")
        rows: list[dict[str, object]] = []
        for values in rows_iter:
            row = {headers[index]: value for index, value in enumerate(values)}
            rows.append(row)
        rows = coerce_rows(rows)
        normalized_path = path.with_suffix(f".{worksheet.title}.parquet")
        _write_parquet(rows, normalized_path)
        return ParsedFile(
            fields=fields_from_rows(rows),
            row_count=len(rows),
            sample_rows=tuple(
                {key: _normalize_value(value) for key, value in row.items()}
                for row in rows[:max_rows]
            ),
            normalized_path=normalized_path,
        )
    finally:
        workbook.close()


def _excel_sheet_names_sync(path: Path) -> tuple[str, ...]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return tuple(workbook.sheetnames)
    finally:
        workbook.close()


async def excel_sheet_names(path: Path) -> tuple[str, ...]:
    try:
        with TemporaryDirectory(prefix=".parse-", dir=path.parent) as temporary:
            stage = Path(temporary).resolve()
            source = stage / path.name
            shutil.copyfile(path, source)
            os.chmod(source, 0o600)
            result = await run_json_worker(
                {"operation": "sheets", "path": str(source)}, directory=stage
            )
            return tuple(result["sheet_names"])
    except (ProcessorFailed, OSError, ValueError, KeyError) as error:
        raise FileParseInvalid("The Excel workbook cannot be decoded safely.") from error


def _parse_json(path: Path, *, max_rows: int) -> ParsedFile:
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as error:
        raise FileParseInvalid("The JSON document is invalid.") from error
    if not isinstance(payload, list) or not payload:
        raise FileParseInvalid("JSON must be a non-empty object array.")
    rows: list[dict[str, object]] = []
    for item in payload:
        if not isinstance(item, dict):
            raise FileParseInvalid("JSON rows must be objects.")
        if any(isinstance(value, (dict, list)) for value in item.values()):
            raise FileParseInvalid("Nested JSON objects are not supported.")
        rows.append(item)
    rows = coerce_rows(rows)
    normalized_path = path.with_suffix(".normalized.parquet")
    _write_parquet(rows, normalized_path)
    return ParsedFile(
        fields=fields_from_rows(rows),
        row_count=len(rows),
        sample_rows=tuple(
            {key: _normalize_value(value) for key, value in row.items()} for row in rows[:max_rows]
        ),
        normalized_path=normalized_path,
    )


def _map_arrow_type(arrow_type: pa.DataType) -> DataType:
    if pa.types.is_boolean(arrow_type):
        return DataType.BOOLEAN
    if pa.types.is_date(arrow_type):
        return DataType.DATE
    if pa.types.is_timestamp(arrow_type):
        return DataType.DATETIME
    if pa.types.is_integer(arrow_type):
        return DataType.INTEGER
    if pa.types.is_floating(arrow_type) or pa.types.is_decimal(arrow_type):
        return DataType.NUMBER
    return DataType.STRING


def _parse_parquet(path: Path, *, max_rows: int) -> ParsedFile:
    table = pq.read_table(path)
    rows = coerce_rows(table.to_pylist())
    fields = fields_from_rows(rows)
    if not rows:
        fields = tuple(
            DatasetField(name=field.name, data_type=_map_arrow_type(field.type))
            for field in table.schema
        )
    sample_rows = tuple(
        {key: _normalize_value(value) for key, value in row.items()}
        for row in rows[:max_rows]
    )
    return ParsedFile(
        fields=fields,
        row_count=table.num_rows,
        sample_rows=sample_rows,
        normalized_path=path,
    )


def _parse_sync(
    path: Path,
    file_format: FileFormat,
    *,
    sheet_name: str | None,
    max_rows: int,
) -> ParsedFile:
    if file_format is FileFormat.CSV:
        return _parse_csv(path, max_rows=max_rows)
    if file_format is FileFormat.EXCEL:
        return _parse_excel(path, sheet_name=sheet_name, max_rows=max_rows)
    if file_format is FileFormat.JSON:
        return _parse_json(path, max_rows=max_rows)
    if file_format is FileFormat.PARQUET:
        return _parse_parquet(path, max_rows=max_rows)
    raise FileParseInvalid("The file format is not supported.")


async def parse_file(
    path: Path,
    file_format: FileFormat,
    *,
    sheet_name: str | None,
    max_rows: int,
) -> ParsedFile:
    if max_rows < 1 or max_rows > 5000:
        raise FileParseInvalid("The file preview row limit is invalid.")
    try:
        _fingerprint(path)  # Reject a symlink before resolving its target.
        path = path.resolve(strict=True)
        key = (path, _fingerprint(path), file_format, sheet_name, max_rows)
        cache = _parse_caches.setdefault(asyncio.get_running_loop(), _ParseCache())
        entry = cache.entries.get(key)
        if entry is not None:
            payload, normalized, fingerprint = entry
            try:
                valid = _fingerprint(normalized) == fingerprint
            except (OSError, FileParseInvalid):
                valid = False
            if valid:
                cache.entries.move_to_end(key)
                return ParsedFile.model_validate_json(payload)
            cache.remove(key)
        flight = cache.flights.get(key)
        if flight is None:
            flight = _ParseFlight(asyncio.create_task(_parse_and_cache(cache, key)))
            cache.flights[key] = flight
        flight.waiters += 1
        try:
            # One cancelled reader must not cancel work still needed by others.
            payload = await asyncio.shield(flight.task)
            if _fingerprint(path) != key[1]:
                raise FileParseInvalid("The source changed during file processing.")
            parsed = ParsedFile.model_validate_json(payload)
            _fingerprint(parsed.normalized_path)
            return parsed
        finally:
            flight.waiters -= 1
            if flight.waiters == 0:
                cache.flights.pop(key, None)
                if not flight.task.done():
                    flight.task.cancel()
                    # Last-reader cancellation must reap the child before the
                    # caller can delete its upload or start another operation.
                    await complete_cleanup(asyncio.gather(flight.task, return_exceptions=True))
    except OSError as error:
        raise FileParseInvalid("The file is unavailable for processing.") from error


async def _parse_file_uncached(
    path: Path,
    file_format: FileFormat,
    *,
    sheet_name: str | None,
    max_rows: int,
    fingerprint: _Fingerprint,
) -> ParsedFile:
    try:
        with TemporaryDirectory(prefix=".parse-", dir=path.parent) as temporary:
            stage = Path(temporary).resolve()
            source = stage / path.name
            shutil.copyfile(path, source)
            os.chmod(source, 0o600)
            result = await run_json_worker(
                {
                    "operation": "file",
                    "path": str(source),
                    "format": file_format.value,
                    "sheet_name": sheet_name,
                    "max_rows": max_rows,
                },
                directory=stage,
            )
            parsed = ParsedFile.model_validate(result["parsed"])
            normalized = parsed.normalized_path
            if normalized.is_symlink() or normalized.parent != stage or not normalized.is_file():
                raise FileParseInvalid("The processor returned an invalid normalized file.")
            if _fingerprint(path) != fingerprint:
                raise FileParseInvalid("The source changed during file processing.")
            destination = path if normalized == source else path.parent / normalized.name
            if normalized != source:
                os.chmod(normalized, 0o600)
                os.replace(normalized, destination)
            return parsed.model_copy(update={"normalized_path": destination})
    except ProcessorBusy as error:
        raise FileParseBusy() from error
    except (ProcessorFailed, OSError, ValueError, KeyError) as error:
        if isinstance(error, FileParseInvalid):
            raise
        raise FileParseInvalid("The file cannot be decoded within processing limits.") from error
